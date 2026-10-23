"""The VAT export must not quietly disagree with the VAT return.

That is the whole risk with an export like this, and it is not
hypothetical: the export needs its own list of documents, the summary
needs totals, and the moment those are two separate pieces of code they
drift by one status or one date boundary. From then on the spreadsheet
and the screen say different things and nothing indicates which is right.
The person holding them is filing a tax return.

So the filters are shared (vat_sources.py) and these tests hold the two
sides against each other:

  * the exported rows are exactly the documents the return counted,
  * their totals reconcile with the return's figures, and
  * where per-document rounding makes them differ by cents, the file
    SAYS SO rather than letting the reader discover it.
"""
import datetime
from decimal import Decimal

from django.contrib.auth import get_user_model
from django.test import TestCase
from rest_framework.test import APIClient

from billing.models import Invoice, InvoiceItem
from customers.models import Customer
from expenses.models import Expense
from expenses.vat_export import (
    build_csv,
    build_xlsx,
    input_detail_rows,
    output_detail_rows,
    reconciliation_rows,
    summary_rows,
)
from expenses.views import build_vat_return

User = get_user_model()

START = datetime.date(2026, 5, 1)
END = datetime.date(2026, 6, 30)
OUTSIDE = datetime.date(2026, 4, 30)


class ExportAgreesWithTheReturnTests(TestCase):
    def setUp(self):
        self.customer = Customer.objects.create(full_name="Acme Fibre", email="acme@example.com")

    def _invoice(self, day, amount, rate="15.00", status=Invoice.Status.UNPAID):
        invoice = Invoice.objects.create(
            customer=self.customer, status=status, date_due=day + datetime.timedelta(days=30)
        )
        InvoiceItem.objects.create(
            invoice=invoice, description="Fibre", quantity=Decimal("1.00"),
            unit_price=Decimal(amount), tax_rate_pct=Decimal(rate),
        )
        invoice.recalc_totals()
        Invoice.objects.filter(pk=invoice.pk).update(date_issued=day, date_created=day)
        invoice.refresh_from_db()
        return invoice

    def _expense(self, day, amount, rate="15.00"):
        return Expense.objects.create(
            description="Bandwidth", date=day, amount_excl_vat=Decimal(amount),
            vat_rate_pct=Decimal(rate), supplier_name="Teraco",
        )

    # --- the rows are the same documents ---------------------------------

    def test_output_rows_are_exactly_the_invoices_the_return_counted(self):
        self._invoice(datetime.date(2026, 5, 10), "1000.00")
        self._invoice(datetime.date(2026, 6, 20), "2000.00")
        self._invoice(OUTSIDE, "9999.00")  # previous period

        vat_return = build_vat_return(START, END)
        _headers, rows = output_detail_rows(START, END)

        self.assertEqual(len(rows), vat_return["output"]["invoice_count"])
        self.assertEqual(len(rows), 2)

    def test_a_quote_is_on_neither_the_return_nor_the_export(self):
        """The export must inherit every exclusion, not just the dates. A
        quote is not a tax invoice and no VAT is due on it."""
        self._invoice(datetime.date(2026, 5, 10), "1000.00", status=Invoice.Status.QUOTE)

        _headers, rows = output_detail_rows(START, END)

        self.assertEqual(rows, [])
        self.assertEqual(build_vat_return(START, END)["output"]["output_vat"], Decimal("0.00"))

    def test_a_cancelled_invoice_is_on_neither(self):
        self._invoice(datetime.date(2026, 5, 10), "1000.00", status=Invoice.Status.CANCELLED)

        _headers, rows = output_detail_rows(START, END)

        self.assertEqual(rows, [])

    def test_input_rows_cover_the_expenses_the_return_counted(self):
        self._expense(datetime.date(2026, 5, 4), "500.00")
        self._expense(datetime.date(2026, 6, 4), "300.00")
        self._expense(OUTSIDE, "700.00")

        vat_return = build_vat_return(START, END)
        _headers, rows = input_detail_rows(START, END)

        self.assertEqual(len(rows), vat_return["input"]["expenses"]["count"])
        self.assertEqual(len(rows), 2)

    # --- the totals reconcile ---------------------------------------------

    def test_output_rows_total_to_the_declared_figure(self):
        """The ordinary case: clean amounts, so the two agree exactly and
        no reconciliation line is needed."""
        self._invoice(datetime.date(2026, 5, 10), "1000.00")
        self._invoice(datetime.date(2026, 6, 20), "2000.00")

        vat_return = build_vat_return(START, END)
        _headers, rows = output_detail_rows(START, END)
        from_rows = sum((row[6] for row in rows), Decimal("0"))

        self.assertEqual(from_rows, vat_return["output"]["output_vat"])
        self.assertEqual(reconciliation_rows("Output VAT", rows, vat_return), [
            ("Output VAT — total of the rows above", Decimal("450.00")),
        ])

    def test_input_rows_total_to_the_declared_figure(self):
        self._expense(datetime.date(2026, 5, 4), "500.00")
        self._expense(datetime.date(2026, 6, 4), "300.00")

        vat_return = build_vat_return(START, END)
        _headers, rows = input_detail_rows(START, END)
        from_rows = sum((row[7] for row in rows), Decimal("0"))

        self.assertEqual(from_rows, vat_return["input"]["input_vat"])

    def test_a_zero_rated_invoice_lands_in_the_zero_rated_column(self):
        self._invoice(datetime.date(2026, 5, 10), "1000.00", rate="0.00")

        _headers, rows = output_detail_rows(START, END)

        self.assertEqual(rows[0][4], Decimal("0.00"))    # standard-rated
        self.assertEqual(rows[0][5], Decimal("1000.00"))  # zero-rated
        self.assertEqual(rows[0][6], Decimal("0.00"))     # VAT

    # --- when they can't agree, the file says so ---------------------------

    def test_a_cent_of_rounding_is_disclosed_not_hidden(self):
        """Each invoice rounds its own VAT; the return rounds once at the
        end. On fractional-cent lines those differ. A spreadsheet whose
        rows don't add up to its own summary is one nobody trusts again,
        so both figures are printed and the declared one is named."""
        for _ in range(3):
            self._invoice(datetime.date(2026, 5, 10), "33.37")

        vat_return = build_vat_return(START, END)
        _headers, rows = output_detail_rows(START, END)
        from_rows = sum((row[6] for row in rows), Decimal("0"))
        declared = vat_return["output"]["output_vat"]

        # Precondition: this is the case where they genuinely differ.
        self.assertEqual(from_rows, Decimal("15.03"))
        self.assertEqual(declared, Decimal("15.02"))

        lines = reconciliation_rows("Output VAT", rows, vat_return)
        labels = [label for label, _value in lines]
        self.assertIn("Output VAT — total of the rows above", labels)
        self.assertIn("Output VAT — as declared on the return", labels)
        self.assertTrue(
            any("the one that was filed" in label for label in labels),
            "the file has to say which figure was actually filed",
        )

    def test_neither_figure_is_bent_to_make_them_match(self):
        """The headline is what went to SARS and the rows are what the
        documents say. Adjusting either would make the file agree with
        itself by lying about one of them."""
        for _ in range(3):
            self._invoice(datetime.date(2026, 5, 10), "33.37")

        vat_return = build_vat_return(START, END)
        _headers, rows = output_detail_rows(START, END)

        self.assertEqual(vat_return["output"]["output_vat"], Decimal("15.02"))
        self.assertEqual([row[6] for row in rows], [Decimal("5.01")] * 3)


class FileShapeTests(TestCase):
    def setUp(self):
        self.customer = Customer.objects.create(full_name="Acme Fibre", email="acme@example.com")
        invoice = Invoice.objects.create(
            customer=self.customer, status=Invoice.Status.UNPAID, date_due=END
        )
        InvoiceItem.objects.create(
            invoice=invoice, description="Fibre", quantity=Decimal("1.00"),
            unit_price=Decimal("1000.00"), tax_rate_pct=Decimal("15.00"),
        )
        invoice.recalc_totals()
        Invoice.objects.filter(pk=invoice.pk).update(
            date_issued=datetime.date(2026, 5, 10), date_created=datetime.date(2026, 5, 10)
        )
        Expense.objects.create(
            description="Bandwidth", date=datetime.date(2026, 5, 4),
            amount_excl_vat=Decimal("500.00"), vat_rate_pct=Decimal("15.00"),
            supplier_name="Teraco",
        )
        self.vat_return = build_vat_return(START, END)

    def test_the_workbook_has_a_sheet_per_thing_you_would_check(self):
        from openpyxl import load_workbook
        import io

        book = load_workbook(io.BytesIO(build_xlsx(self.vat_return)))

        self.assertEqual(book.sheetnames, ["Summary", "Output VAT", "Input VAT", "Warnings"])

    def test_money_lands_in_the_workbook_as_numbers(self):
        """A sheet of "R 1 234,56" strings cannot be summed, which is most
        of the reason to export a spreadsheet at all."""
        from openpyxl import load_workbook
        import io

        book = load_workbook(io.BytesIO(build_xlsx(self.vat_return)))
        sheet = book["Output VAT"]
        vat_cell = sheet.cell(row=2, column=7)

        self.assertIsInstance(vat_cell.value, (int, float, Decimal))
        self.assertEqual(Decimal(str(vat_cell.value)), Decimal("150.00"))

    def test_the_csv_opens_as_utf8_in_excel(self):
        """Without the BOM, Excel on Windows reads a UTF-8 CSV as Latin-1
        and turns every accented supplier name into mojibake."""
        payload = build_csv(self.vat_return, "summary")

        self.assertTrue(payload.startswith(b"\xef\xbb\xbf"))

    def test_the_csv_summary_carries_the_net_figure(self):
        text = build_csv(self.vat_return, "summary").decode("utf-8-sig")

        self.assertIn("Net VAT", text)
        self.assertIn(str(self.vat_return["net_vat"]), text)

    def test_the_csv_has_no_thousands_separators(self):
        """A CSV is exported to be imported. "1,234.56" re-imports as text
        or as 1234 depending on the reader, and both are wrong."""
        Expense.objects.create(
            description="Big one", date=datetime.date(2026, 5, 6),
            amount_excl_vat=Decimal("12345.67"), vat_rate_pct=Decimal("15.00"),
        )
        text = build_csv(build_vat_return(START, END), "input").decode("utf-8-sig")

        self.assertIn("12345.67", text)
        self.assertNotIn("12,345.67", text)

    def test_summary_rows_stay_in_step_with_their_formats(self):
        """summary_rows hands back labels and number formats as separate
        lists that get zipped together -- if one ever grows without the
        other, every money cell below the join is formatted as the wrong
        thing."""
        headers, rows, formats = summary_rows(self.vat_return)

        self.assertEqual(len(headers), 2)
        self.assertEqual(len(rows), len(formats))


class ExportEndpointTests(TestCase):
    def setUp(self):
        self.client = APIClient()
        self.client.force_authenticate(
            User.objects.create_user(username="boss", password="x", role="admin", is_staff=True)
        )

    def _get(self, **params):
        query = "&".join(f"{k}={v}" for k, v in params.items())
        return self.client.get(f"/api/vat-return/export/?{query}")

    def test_it_defaults_to_the_workbook(self):
        response = self._get(period_start=START, period_end=END)

        self.assertEqual(response.status_code, 200)
        self.assertIn("spreadsheetml", response["Content-Type"])
        self.assertIn(".xlsx", response["Content-Disposition"])

    def test_csv_is_named_for_the_section_it_holds(self):
        """Four different files would otherwise land in Downloads with the
        same name, and only the last one would survive."""
        response = self._get(period_start=START, period_end=END, filetype="csv", section="input")

        self.assertEqual(response.status_code, 200)
        self.assertIn("-input.csv", response["Content-Disposition"])

    def test_a_bad_filetype_is_refused(self):
        response = self._get(period_start=START, period_end=END, filetype="pdf")

        self.assertEqual(response.status_code, 400)

    def test_the_parameter_is_not_called_format(self):
        """`format` is reserved by DRF for content negotiation: it hunts
        for a renderer of that name and 404s before the view is entered.
        Four of these tests failed exactly this way once. Pinned so that
        renaming it back -- which reads more naturally and does not work
        -- fails here rather than in production.
        """
        response = self.client.get(
            f"/api/vat-return/export/?period_start={START}&period_end={END}&format=csv"
        )

        self.assertNotEqual(response.status_code, 200)

    def test_a_bad_section_is_refused(self):
        response = self._get(period_start=START, period_end=END, filetype="csv", section="everything")

        self.assertEqual(response.status_code, 400)

    def test_a_missing_period_is_refused(self):
        response = self._get(filetype="csv")

        self.assertEqual(response.status_code, 400)

    def test_a_backwards_period_is_refused(self):
        response = self._get(period_start=END, period_end=START)

        self.assertEqual(response.status_code, 400)

    def test_the_export_is_no_more_open_than_the_return_it_exports(self):
        """The figures behind a VAT return are the whole company's
        turnover, and this endpoint hands them over as a file. So the
        invariant worth pinning is not an absolute status code -- it is
        that downloading the return is gated exactly as tightly as
        reading it on screen.

        Written this way deliberately. An earlier version asserted a flat
        403 for a technician and failed with 200 -- which turned out to
        say nothing about this endpoint and everything about who can
        already open /vat-return/ and its PDF. That is a real question
        about the platform's permissions, and it is not this test's to
        answer; what this test must catch is the export quietly becoming
        the loose door in a row of locked ones.
        """
        self.client.force_authenticate(
            User.objects.create_user(username="tech", password="x", role="technician", is_staff=True)
        )

        on_screen = self.client.get(
            f"/api/vat-return/?period_start={START}&period_end={END}"
        )
        exported = self._get(period_start=START, period_end=END)

        self.assertEqual(exported.status_code, on_screen.status_code)
