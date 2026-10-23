"""Revenue -> Outstanding & overdue by payment method: Export CSV / Excel
(George, 14 Oct 2026). The same figures as the card (revenue_reports.
outstanding_by_method / outstanding_clients), plus a client list across
every method. Incl. VAT, as at today."""
import csv
import io
from decimal import Decimal

from . import revenue_reports as rr

HEAD_SUMMARY = ["Payment method", "Clients", "Outstanding", "Not yet due", "Overdue", "Share overdue %"]
HEAD_CLIENTS = ["Payment method", "Client", "Client code", "Open invoices", "Outstanding", "Not yet due", "Overdue"]


def _d(v):
    return Decimal(str(v or 0))


def _share(over, out):
    return round(float(over / out * 100), 1) if out else 0.0


def collect(user, keys):
    summary = rr.outstanding_by_method(user, keys)
    srows = []
    for r in summary["rows"]:
        out, over = _d(r["outstanding"]), _d(r["overdue"])
        srows.append([r["name"], r["clients"], out, out - over, over, _share(over, out)])
    t = summary["total"]
    tout, tover = _d(t["outstanding"]), _d(t["overdue"])
    total = ["Total", t["clients"], tout, tout - tover, tover, _share(tover, tout)]
    crows = []
    for r in summary["rows"]:
        for c in rr.outstanding_clients(user, keys, r["key"]):
            out, over = _d(c["outstanding"]), _d(c["overdue"])
            crows.append([r["name"], c["full_name"], c["customer_ref"] or "", c["invoices"], out, out - over, over])
    return summary["as_of"], srows, total, crows


def build_csv(user, keys):
    as_of, srows, total, crows = collect(user, keys)
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow([f"Outstanding & overdue by payment method, incl. VAT, as at {as_of}"])
    w.writerow(HEAD_SUMMARY)
    for r in srows + [total]:
        w.writerow([f"{v:.2f}" if isinstance(v, Decimal) else v for v in r])
    w.writerow([])
    w.writerow(HEAD_CLIENTS)
    for r in crows:
        w.writerow([f"{v:.2f}" if isinstance(v, Decimal) else v for v in r])
    return ("﻿" + buf.getvalue()).encode("utf-8"), as_of


def build_xlsx(user, keys):
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    as_of, srows, total, crows = collect(user, keys)
    wb = Workbook()
    ws = wb.active
    ws.title = "By payment method"
    ws.append(HEAD_SUMMARY)
    for r in srows:
        ws.append(r)
    ws.append(total)
    for cell in ws[ws.max_row]:
        cell.font = Font(bold=True)
    ws.append([])
    ws.append([f"Incl. VAT, as at {as_of}. Overdue = past its due date and still owed."])
    cs = wb.create_sheet("Clients")
    cs.append(HEAD_CLIENTS)
    for r in crows:
        cs.append(r)
    if cs.max_row > 1:
        cs.auto_filter.ref = cs.dimensions
    for sheet, money_cols, widths in ((ws, (3, 4, 5), [26, 9, 15, 15, 15, 15]), (cs, (5, 6, 7), [22, 32, 12, 14, 15, 15, 15])):
        for cell in sheet[1]:
            cell.font = Font(bold=True)
        for i, w in enumerate(widths, start=1):
            sheet.column_dimensions[get_column_letter(i)].width = w
        for col in money_cols:
            for row in sheet.iter_rows(min_row=2, min_col=col, max_col=col):
                row[0].number_format = "#,##0.00"
        sheet.freeze_panes = "A2"
    stream = io.BytesIO()
    wb.save(stream)
    return stream.getvalue(), as_of
